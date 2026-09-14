import re
from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}

# Matches "Point 1: promotions", "Point 2: Account Management", etc.
# After normalization the whole doc is one line, so this is reliable.
POINT_RE = re.compile(r"Point\s+(\d+)\s*:\s*([^:]+?)(?=\s+\d+\)\s)", re.IGNORECASE)

# Matches question starts like "1) How ...", "14) What ...".
# Answers use "1." style lists, so "N)" is (almost) always a new question.
QUESTION_RE = re.compile(r"(?<!\d)(\d{1,3})\)\s+")

# If a single Q+A is longer than this, sub-split the answer only
# (keeping the question header on every sub-chunk).
MAX_QA_CHARS = 1500


# load the knowledge based file and return a list of document objects
def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def normalize_text(text: str) -> str:
    """Fix the 'one word per line' extraction that pypdf/PyPDFLoader
    produces for this PDF, plus ligatures / stray whitespace."""
    text = text.replace("ﬁ", "fi").replace("ﬂ", "fl")  # \ufb01, \ufb02
    text = text.replace("\u2019", "'").replace("\u2018", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = re.sub(r"\s+", " ", text)  # newlines, double spaces -> single space
    # fix spacing around punctuation broken by extraction
    text = re.sub(r"\s+([?.!,;:])", r"\1", text)
    text = re.sub(r"\(\s+", "(", text)
    text = re.sub(r"\s+\)", ")", text)
    return text.strip()


def _split_questions(point_body: str):
    """Yield (question_number, question_block_text) inside one Point section."""
    starts = list(QUESTION_RE.finditer(point_body))
    for i, m in enumerate(starts):
        end = starts[i + 1].start() if i + 1 < len(starts) else len(point_body)
        block = point_body[m.start():end].strip()
        if len(block) < 10:  # skip noise
            continue
        yield m.group(1), block


def _build_qa_chunk(
    *, question_number: str, block: str,
    point_number: str, point_title: str, source: str,
) -> tuple[str, str, str]:
    """Split a 'N) Question? Answer...' block into (question, answer, page_content).

    page_content carries a context header so the chunk is self-contained
    when retrieved alone from the vector DB (contextual retrieval).
    """
    # block starts with "N) ". Strip it, then question = up to first "?"
    body = QUESTION_RE.sub("", block, count=1).strip()
    q_end = body.find("?")
    if q_end != -1 and q_end < 300:
        question = body[: q_end + 1].strip()
        answer = body[q_end + 1 :].strip()
    else:  # no "?" found (malformed) -> first sentence is the question
        parts = re.split(r"(?<=[.?])\s+", body, maxsplit=1)
        question = parts[0].strip()
        answer = parts[1].strip() if len(parts) > 1 else ""
    if not answer:
        answer = "(No answer text extracted)"
    header = f"[Point {point_number}: {point_title} | Q{question_number}: {question}]"
    page_content = f"{header}\nQuestion: {question}\nAnswer: {answer}"
    return question, answer, page_content


def chunk_qa_document(full_text: str, source: str) -> list[Document]:
    """Structure-aware split: 1 question + its answer = 1 chunk."""
    text = normalize_text(full_text)
    point_matches = list(POINT_RE.finditer(text))
    if not point_matches:
        return []  # caller falls back to generic splitting

    chunks: list[Document] = []
    fallback_splitter = RecursiveCharacterTextSplitter(
        chunk_size=800, chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    for pi, pm in enumerate(point_matches):
        point_number = pm.group(1)
        point_title = re.sub(r"\s+", " ", pm.group(2)).strip()
        section_end = (
            point_matches[pi + 1].start() if pi + 1 < len(point_matches) else len(text)
        )
        point_body = text[pm.end():section_end]

        for q_num, block in _split_questions(point_body):
            question, answer, page_content = _build_qa_chunk(
                question_number=q_num, block=block,
                point_number=point_number, point_title=point_title,
                source=source,
            )
            base_meta = {
                "source": source,
                "point_number": point_number,
                "point_title": point_title,
                "question_number": q_num,
                "question": question,
                "chunk_type": "qa_pair",
            }
            # Long answers (e.g. return policies) -> sub-split answer only,
            # repeating the Q header so every sub-chunk stays retrievable.
            if len(page_content) > MAX_QA_CHARS:
                sub_docs = fallback_splitter.split_text(answer)
                for si, sub in enumerate(sub_docs):
                    chunks.append(Document(
                        page_content=(
                            f"[Point {point_number}: {point_title} "
                            f"| Q{q_num} (part {si+1}/{len(sub_docs)}): {question}]\n"
                            f"Question: {question}\nAnswer: {sub}"
                        ),
                        metadata={**base_meta, "chunk_type": "qa_pair_part",
                                  "part": si + 1, "parts": len(sub_docs)},
                    ))
            else:
                chunks.append(Document(page_content=page_content, metadata=base_meta))
    return chunks


def chunk_generic(docs: Iterable[Document]) -> list[Document]:
    """Fallback for non-QA docs (.txt/.md/generic): fixed-size with overlap."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800, chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_documents(list(docs))


# split the documents into smaller chunks
def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    docs = list(docs)
    if not docs:
        return []
    source = docs[0].metadata.get("source", "")

    # IMPORTANT: join ALL pages first. A Q+A often spans a page boundary
    # (e.g. Q4 starts on p0, answer continues on p1). Splitting per-page
    # would cut answers in half.
    full_text = " ".join(d.page_content for d in docs)

    qa_chunks = chunk_qa_document(full_text, source)
    if qa_chunks:
        return qa_chunks

    # Not a Point/Q&A doc -> generic character splitting
    return chunk_generic(docs)